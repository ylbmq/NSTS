"""Create a reproducible Excel workbook for manual validation of benchmark labels."""

import argparse
import csv
import hashlib
import json
import random
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


DATA_DIR = Path(__file__).resolve().parent / "data"
WORKBOOK_FILENAME = "label_validation.xlsx"
DATE_FORMAT = "yyyy-mm-dd"
FIELDNAMES = ["url", "automatic_timestamp", "manual_timestamp", "label_correctness"]


def excel_date(value):
    """Convert an ISO date or existing Python date to an Excel date value."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(value.strip())


def add_validation_sheet(workbook: Workbook, domain: str, rows: list[dict]) -> None:
    """Add one domain's audit sheet, preserving supplied manual dates."""
    sheet = workbook.create_sheet(title=domain)
    sheet.append(FIELDNAMES)
    sheet.freeze_panes = "C2"
    sheet.column_dimensions["A"].width = 100
    sheet.column_dimensions["B"].width = 23
    sheet.column_dimensions["B"].hidden = True
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 22
    for column in ("B", "C"):
        sheet.column_dimensions[column].number_format = DATE_FORMAT
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, row in enumerate(rows, 2):
        url = row["url"]
        # Older CSVs store HYPERLINK formulas; retain those when converting.
        sheet.cell(index, 1, url)
        if not url.startswith("=HYPERLINK("):
            sheet.cell(index, 1).hyperlink = url
        sheet.cell(index, 1).style = "Hyperlink"
        sheet.cell(index, 2, excel_date(row["automatic_timestamp"])).number_format = DATE_FORMAT
        manual = sheet.cell(index, 3, excel_date(row.get("manual_timestamp")))
        manual.number_format = DATE_FORMAT
        sheet.cell(index, 4,
            f'=IF(TRIM(C{index}&"")="","",IFERROR(IF(INT(B{index})='
            f'IF(ISNUMBER(C{index}),INT(C{index}),DATEVALUE(TRIM(C{index}))),'
            f'"Correct","Incorrect"),"Incorrect"))')
    last = len(rows) + 1
    # Excel formats recognized date entries automatically. Validation rejects
    # unrecognized typed text; pasting can bypass Excel's validation rules.
    date_entry = DataValidation(
        type="date", operator="between", formula1="DATE(1900,1,1)",
        formula2="DATE(9999,12,31)", allow_blank=True,
    )
    date_entry.showInputMessage = True
    date_entry.promptTitle = "Enter a date"
    date_entry.prompt = "Enter a date such as 2025-01-31. Dates display as yyyy-mm-dd."
    date_entry.showErrorMessage = True
    date_entry.errorStyle = "stop"
    date_entry.errorTitle = "Valid date required"
    date_entry.error = "Enter an Excel-recognized date, preferably yyyy-mm-dd, or leave blank."
    sheet.add_data_validation(date_entry)
    date_entry.add(f"B2:C{last}")
    results = f"D2:D{last}"
    sheet.auto_filter.ref = f"A1:D{last}"
    for result, color in [("Correct", "C6EFCE"), ("Incorrect", "FFC7CE")]:
        sheet.conditional_formatting.add(results, FormulaRule(
            formula=[f'D2="{result}"'], fill=PatternFill("solid", fgColor=color)))
    reviewed = f'COUNTIF({results},"Correct")+COUNTIF({results},"Incorrect")'
    sheet.cell(last + 1, 3, "Label accuracy (reviewed)")
    sheet.cell(last + 1, 4, f'=IFERROR(COUNTIF({results},"Correct")/({reviewed}),"")')
    sheet.cell(last + 1, 4).number_format = "0.00%"
    sheet.cell(last + 2, 3, "Reviewed / sampled")
    sheet.cell(last + 2, 4, f'=({reviewed})&" / {len(rows)}"')


def add_summary_sheet(workbook: Workbook, sample_sizes: dict[str, int]) -> None:
    """Link each domain's reviewed accuracy, leaving unreviewed domains blank."""
    sheet = workbook.create_sheet("Summary")
    sheet.append(["domain", "label_accuracy"])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 24
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, (domain, size) in enumerate(sample_sizes.items(), 2):
        sheet.cell(index, 1, domain)
        reference = "'" + domain.replace("'", "''") + f"'!D{size + 2}"
        sheet.cell(index, 2, f'=IF({reference}="","",{reference})')
        sheet.cell(index, 2).number_format = "0.00%"
    sheet.auto_filter.ref = f"A1:B{len(sample_sizes) + 1}"


def write_workbook(path: Path, samples: dict[str, list[dict]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for domain, rows in samples.items():
        add_validation_sheet(workbook, domain, rows)
    add_summary_sheet(workbook, {domain: len(rows) for domain, rows in samples.items()})
    with path.open("xb") as handle:
        workbook.save(handle)


def convert_csvs(output_dir: Path) -> None:
    """Convert existing samples without resampling or changing CSV annotations."""
    files = sorted(output_dir.glob("*_validation.csv"))
    if not files:
        raise ValueError(f"No validation CSVs found in {output_dir}")
    target = output_dir / WORKBOOK_FILENAME
    if target.exists():
        raise FileExistsError(f"Refusing to overwrite {target}")
    samples = {}
    for path in files:
        with path.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        if not rows:
            raise ValueError(f"Empty sample: {path}")
        samples[path.stem.removesuffix("_validation")] = rows
    write_workbook(target, samples)
    print(f"Created {target}: {len(samples)} domain tabs + Summary")


def generate_samples(label_dir: Path, output_dir: Path, sample_size: int, seed: int) -> None:
    if sample_size < 1:
        raise ValueError("Sample size must be positive.")
    label_files = sorted(label_dir.glob("*_labels.json"))
    if not label_files:
        raise ValueError(f"No *_labels.json files found in {label_dir}")

    # Prepare every sample before writing anything, including checking for duplicates.
    samples = {}
    manifest = {"seed": seed, "requested_sample_size": sample_size, "domains": {}}
    for label_file in label_files:
        domain = label_file.stem.removesuffix("_labels")
        contents = label_file.read_bytes()
        labels = json.loads(contents)
        if not labels:
            raise ValueError(f"Empty label file: {label_file}")
        urls = [row["url"] for row in labels]
        if len(urls) != len(set(urls)):
            raise ValueError(f"Duplicate URLs in {label_file}; resolve before sampling.")
        for row in labels:
            if not row["url"] or not row["timestamp"]:
                raise ValueError(f"Missing URL or timestamp in {label_file}")

        # Domain-specific RNG keeps each sample independent of other domains.
        rng = random.Random(f"{seed}:{domain}")
        sample = rng.sample(sorted(labels, key=lambda row: row["url"]), min(sample_size, len(labels)))
        samples[domain] = sample
        manifest["domains"][domain] = {
            "label_file": label_file.name,
            "label_sha256": hashlib.sha256(contents).hexdigest(),
            "population_size": len(labels),
            "sample_size": len(sample),
            "workbook_file": WORKBOOK_FILENAME,
            "sheet_name": domain,
        }

    targets = [output_dir / WORKBOOK_FILENAME, output_dir / "sampling_manifest.json"]
    existing = [str(path) for path in targets if path.exists()]
    if existing:
        raise FileExistsError(
            "Refusing to overwrite existing validation files. Use a new --output-dir.\n"
            + "\n".join(existing)
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    write_workbook(targets[0], {
        domain: [{"url": row["url"], "automatic_timestamp": row["timestamp"]}
                 for row in sample]
        for domain, sample in samples.items()
    })
    print(f"Created {targets[0]}: {len(samples)} domain tabs + Summary")
    with targets[-1].open("x", encoding="utf-8") as handle:
        json.dump(manifest, handle, indent=2)
        handle.write("\n")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--label-dir", type=Path, default=DATA_DIR / "labels")
    parser.add_argument("--output-dir", type=Path, default=DATA_DIR / "labels" / "validation")
    parser.add_argument("--sample-size", type=int, default=100)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--from-csv", action="store_true",
                        help="Convert existing validation CSVs in --output-dir to Excel without resampling.")
    args = parser.parse_args()
    try:
        if args.from_csv:
            convert_csvs(args.output_dir)
        else:
            generate_samples(args.label_dir, args.output_dir, args.sample_size, args.seed)
    except (ValueError, OSError, KeyError, TypeError) as error:
        parser.exit(1, f"Error: {error}\n")


if __name__ == "__main__":
    main()
